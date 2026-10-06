"""
ỨNG DỤNG WEB TRÍCH XUẤT SỐ SM & NGÀY (NHỰA TIỀN PHONG)
- GIAO DIỆN LỘT XÁC TOÀN DIỆN (LUXURY CYBER TECH / GLASSMORPHISM THEME)
- Tên file Excel cố định: PDF-TO-EXCEL.xlsx
- Tự động download và tự động mở file Excel (khi chạy local).
- Tự động reset về ban đầu khi thêm hoặc xóa bớt file PDF đã chọn.
- Đếm ngược thời gian (phút & giây).
"""

import os
import sys
import re
import io
import time
import base64
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import pypdfium2 as pdfium
import pytesseract
from PIL import Image
import pandas as pd
import streamlit as st
import streamlit.components.v1 as components

# ==================== CẤU HÌNH TESSERACT OCR ====================
def setup_tesseract():
    if os.name == 'nt':
        default_paths = [
            os.path.join(os.path.dirname(os.path.abspath(__file__)), "Tesseract-OCR", "tesseract.exe"),
            r"C:\Program Files\Tesseract-OCR\tesseract.exe",
            r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
            os.path.expanduser(r"~\AppData\Local\Programs\Tesseract-OCR\tesseract.exe")
        ]
        for path in default_paths:
            if os.path.exists(path):
                pytesseract.pytesseract.tesseract_cmd = path
                return

setup_tesseract()

# ==================== HÀM TIỆN ÍCH ====================
def format_time_str(total_seconds):
    """Định dạng thời gian ra phút và giây rõ ràng"""
    sec = max(int(total_seconds), 0)
    mins = sec // 60
    secs = sec % 60
    if mins > 0:
        return f"{mins} phút {secs:02d}s"
    else:
        return f"0 phút {secs:02d}s"

def sanitize_sheet_name(name, existing_names):
    """Quy chuẩn tên Sheet trong Excel (tối đa 31 ký tự, loại bỏ ký tự cấm)"""
    if name.lower().endswith('.pdf'):
        name = name[:-4]
    clean_name = re.sub(r'[\/\\\?\*\[\]\:]', '_', name).strip()
    if not clean_name:
        clean_name = "Sheet"
    base_name = clean_name[:31]
    final_name = base_name
    counter = 1
    while final_name.lower() in existing_names:
        suffix = f"_{counter}"
        max_len = 31 - len(suffix)
        final_name = base_name[:max_len] + suffix
        counter += 1
    existing_names.add(final_name.lower())
    return final_name

def extract_from_single_pdf(file_bytes, page_callback=None):
    """Trích xuất SM và Ngày từ 1 file PDF"""
    pdf = pdfium.PdfDocument(file_bytes)
    total_pages = len(pdf)
    records = []
    
    for page_idx in range(total_pages):
        page_num = page_idx + 1
        page = pdf[page_idx]
        bitmap = page.render(scale=2.0)
        img = bitmap.to_pil()
        w, h = img.size
        
        crop = img.crop((0, 0, w, int(h * 0.30)))
        try:
            text = pytesseract.image_to_string(crop, lang='eng')
        except Exception:
            text = ""
            
        text_u = text.upper()
        is_tp = ('TIEN PHONG' in text_u or 'THIEU NIEN' in text_u or 'NHUATIENPHONG' in text_u)
        
        if is_tp:
            sm_match = re.search(r'[\$S§s]M[\s\.:,]*([0-9]{4})[\.,\s_]*([0-9]{4})', text)
            if not sm_match:
                sm_match = re.search(r'[\$S§s]M[\s\.:,]*([0-9]{4,8})', text)
                val = sm_match.group(1) if sm_match else ""
                sm = f"SM{val[:4]}.{val[4:]}" if len(val) == 8 else (f"SM{val}" if val else None)
            else:
                sm = f"SM{sm_match.group(1)}.{sm_match.group(2)}"
                
            date_match = re.search(r'Ng[aàeè]y[\s;:,-]*([0-9]{1,2})[\/\.\-]([0-9]{1,2})[\/\.\-]([0-9]{4})', text, re.IGNORECASE)
            if not date_match:
                date_match = re.search(r'([0-9]{1,2})[\/\.\-]([0-9]{1,2})[\/\.\-]([0-9]{4})', text)
                
            if date_match:
                d_val = int(date_match.group(1))
                m_val = int(date_match.group(2))
                y_val = date_match.group(3)
                dt = f"{d_val:02d}/{m_val:02d}/{y_val}"
            else:
                dt = None
                
            if sm and dt:
                records.append({
                    'sm': sm,
                    'date': dt,
                    'page': page_num
                })
                
        if page_callback:
            page_callback(page_num, total_pages)
            
    return records

def create_excel_multi_sheet(file_results):
    """Tạo file Excel với mỗi file PDF là 1 sheet riêng"""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    
    fill_head = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    font_head = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    fill_alt = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    border = Border(
        left=Side(style="thin", color="CBD5E1"), right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"), bottom=Side(style="thin", color="CBD5E1")
    )
    headers = ["STT", "SỐ SM", "NGÀY", "SỐ TRANG CỦA SỐ SM"]
    
    existing_sheets = set()
    for file_name, records in file_results:
        sheet_title = sanitize_sheet_name(file_name, existing_sheets)
        ws = wb.create_sheet(title=sheet_title)
        ws.views.sheetView[0].showGridLines = True
        
        ws.row_dimensions[1].height = 26
        for idx, h in enumerate(headers, 1):
            c = ws.cell(row=1, column=idx, value=h)
            c.fill = fill_head
            c.font = font_head
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = border
            
        for r_idx, r in enumerate(records, start=2):
            stt = r_idx - 1
            ws.row_dimensions[r_idx].height = 20
            row_vals = [stt, r['sm'], r['date'], r['page']]
            for c_idx, v in enumerate(row_vals, 1):
                c = ws.cell(row=r_idx, column=c_idx, value=v)
                c.font = Font(name="Calibri", size=11)
                c.border = border
                c.alignment = Alignment(horizontal="center", vertical="center")
                if r_idx % 2 == 1:
                    c.fill = fill_alt
                    
        ws.freeze_panes = "A2"
        ws.column_dimensions['A'].width = 10
        ws.column_dimensions['B'].width = 18
        ws.column_dimensions['C'].width = 16
        ws.column_dimensions['D'].width = 25
        
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()

def auto_open_local_file(file_bytes, filename):
    """Mở file trên Windows khi chạy local"""
    if os.name == 'nt' and hasattr(os, 'startfile'):
        try:
            downloads_dir = os.path.join(os.path.expanduser("~"), "Downloads")
            if not os.path.exists(downloads_dir):
                downloads_dir = os.getcwd()
            target_path = os.path.join(downloads_dir, filename)
            with open(target_path, "wb") as f:
                f.write(file_bytes)
            os.startfile(target_path)
            return target_path
        except Exception:
            pass
    return None

def render_auto_downloader(file_bytes, filename):
    """Tự động kích hoạt tải file về máy qua JavaScript"""
    b64 = base64.b64encode(file_bytes).decode()
    js_code = f"""
    <script>
        function triggerDownload() {{
            try {{
                var targetDoc = (window.parent && window.parent.document) ? window.parent.document : document;
                var a = targetDoc.createElement('a');
                a.href = 'data:application/vnd.openxmlformats-officedocument.spreadsheetml.sheet;base64,{b64}';
                a.download = '{filename}';
                targetDoc.body.appendChild(a);
                a.click();
                setTimeout(function() {{
                    try {{ targetDoc.body.removeChild(a); }} catch(err) {{}}
                }}, 1000);
            }} catch(e) {{
                var a2 = document.createElement('a');
                a2.href = 'data:application/vnd.openxmlformats-officedocument.spreadsheetml.sheet;base64,{b64}';
                a2.download = '{filename}';
                document.body.appendChild(a2);
                a2.click();
            }}
        }}
        setTimeout(triggerDownload, 300);
    </script>
    """
    components.html(js_code, height=0, width=0)

# ==================== GIAO DIỆN CYBER TECH HIỆN ĐẠI ====================
st.set_page_config(
    page_title="AI Scan • Trích Xuất Tiền Phong", 
    page_icon="⚡", 
    layout="centered"
)

# Can thiệp sâu CSS: Nền không gian tối sang trọng (Dark Luxury Tech)
st.markdown("""
<style>
    /* Ẩn triệt để header, menu, footer Streamlit */
    #MainMenu {visibility: hidden;}
    footer {visibility: hidden;}
    header {visibility: hidden;}
    
    /* Nền toàn trang web: Dark Tech Gradient */
    .stApp {
        background: radial-gradient(circle at 50% 0%, #1E293B 0%, #0F172A 70%, #020617 100%) !important;
        color: #F8FAFC !important;
    }
    
    /* Thu gọn đỉnh trang sát mép trên */
    .block-container {
        padding-top: 1rem !important;
        padding-bottom: 0.5rem !important;
        max-width: 820px !important;
    }
    
    /* Tiêu đề Cyber Neon Gradient */
    .cyber-title {
        font-family: 'Segoe UI', -apple-system, sans-serif;
        font-size: 26px;
        font-weight: 900;
        background: linear-gradient(135deg, #38BDF8 0%, #818CF8 50%, #C084FC 100%);
        -webkit-background-clip: text;
        -webkit-text-fill-color: transparent;
        text-align: center;
        margin: 0 0 4px 0;
        text-shadow: 0 0 30px rgba(56, 189, 248, 0.2);
    }
    .cyber-sub {
        font-size: 13px;
        color: #94A3B8;
        text-align: center;
        margin-bottom: 12px;
        font-weight: 500;
    }
    
    /* Thiết kế lại hộp Kéo thả File (Dropzone) hoàn toàn mới */
    [data-testid="stFileUploader"] {
        background: rgba(30, 41, 59, 0.7) !important;
        border: 2px dashed #38BDF8 !important;
        border-radius: 16px !important;
        padding: 10px !important;
        box-shadow: 0 8px 32px rgba(0, 0, 0, 0.3) !important;
        backdrop-filter: blur(10px) !important;
    }
    [data-testid="stFileUploader"] section {
        background: transparent !important;
    }
    [data-testid="stFileUploader"] label {
        color: #E2E8F0 !important;
        font-weight: 600 !important;
    }
    
    /* Thiết kế lại nút BẮT ĐẦU TRÍCH XUẤT: Neon Glowing Button */
    .stButton > button {
        background: linear-gradient(135deg, #0284C7 0%, #2563EB 50%, #4F46E5 100%) !important;
        color: #FFFFFF !important;
        border: 1px solid rgba(255, 255, 255, 0.2) !important;
        border-radius: 14px !important;
        font-size: 16px !important;
        font-weight: 800 !important;
        padding: 12px 24px !important;
        box-shadow: 0 0 25px rgba(37, 99, 235, 0.6) !important;
        transition: all 0.3s ease !important;
        letter-spacing: 0.5px !important;
    }
    .stButton > button:hover {
        transform: translateY(-2px) !important;
        box-shadow: 0 0 35px rgba(56, 189, 248, 0.8) !important;
    }
    
    /* Metric boxes dạng bảng điều khiển phi thuyền hiện đại */
    .glow-card-file {
        background: rgba(30, 41, 59, 0.85);
        border: 1.5px solid #F59E0B;
        border-radius: 12px;
        padding: 10px 12px;
        text-align: center;
        box-shadow: 0 0 15px rgba(245, 158, 11, 0.2);
    }
    .glow-card-prog {
        background: rgba(30, 41, 59, 0.85);
        border: 1.5px solid #10B981;
        border-radius: 12px;
        padding: 10px 12px;
        text-align: center;
        box-shadow: 0 0 15px rgba(16, 185, 129, 0.2);
    }
    .glow-card-time {
        background: rgba(30, 41, 59, 0.85);
        border: 1.5px solid #38BDF8;
        border-radius: 12px;
        padding: 10px 12px;
        text-align: center;
        box-shadow: 0 0 15px rgba(56, 189, 248, 0.2);
    }
    
    .glow-val {
        font-size: 18px;
        font-weight: 800;
        font-family: 'Consolas', monospace;
    }
    .glow-label {
        font-size: 11px;
        color: #94A3B8;
        font-weight: 600;
        margin-top: 2px;
    }
    
    /* Khung kết quả thành công sang trọng */
    .cyber-success {
        background: rgba(6, 78, 59, 0.4);
        border: 2px solid #10B981;
        border-radius: 14px;
        padding: 16px 20px;
        color: #ECFDF5;
        margin-top: 8px;
        margin-bottom: 12px;
        box-shadow: 0 0 30px rgba(16, 185, 129, 0.25);
        backdrop-filter: blur(12px);
    }
</style>
""", unsafe_allow_html=True)

# Header
st.markdown("""
<div style="text-align: center; margin-bottom: 8px;">
    <h1 class="cyber-title">⚡ TRÍCH XUẤT SỐ SM & NGÀY PHIẾU GIAO HÀNG</h1>
    <p class="cyber-sub">AI Scan Nhựa Tiền Phong • Tự động nhận diện & đếm trang • Xuất Excel đa sheet</p>
</div>
""", unsafe_allow_html=True)

# Khung Upload
uploaded_files = st.file_uploader(
    "Kéo thả hoặc duyệt chọn các file PDF scan:", 
    type=["pdf"], 
    accept_multiple_files=True,
    help="Có thể chọn nhiều file cùng lúc. Bấm dấu ✖ bên cạnh file để xóa nếu chọn nhầm."
)

# ==================== TỰ ĐỘNG RESET NẾU THAY ĐỔI FILE ====================
current_filenames = [f.name for f in uploaded_files] if uploaded_files else []
last_filenames = st.session_state.get('last_uploaded_files', None)

if last_filenames is not None and current_filenames != last_filenames:
    for k in ['excel_data', 'excel_name', 'total_files', 'total_extracted', 'total_time_str', 'completed', 'opened_path']:
        if k in st.session_state:
            del st.session_state[k]

st.session_state['last_uploaded_files'] = current_filenames

# ==================== XỬ LÝ CHÍNH ====================
if uploaded_files:
    total_files = len(uploaded_files)
    st.markdown(f"<div style='color: #38BDF8; font-size: 13px; font-weight: 600; margin-bottom: 6px;'>📁 Đang sẵn sàng: {total_files} file PDF</div>", unsafe_allow_html=True)
    
    # 1. NÚT "BẮT ĐẦU TRÍCH XUẤT" (Chỉ hiện khi CHƯA xử lý xong)
    if 'completed' not in st.session_state:
        if st.button("⚡ BẮT ĐẦU TRÍCH XUẤT NGAY", type="primary", use_container_width=True):
            
            with st.spinner("Đang chuẩn bị quét dữ liệu..."):
                pages_per_file = []
                for f in uploaded_files:
                    try:
                        pdf_temp = pdfium.PdfDocument(f.getvalue())
                        pages_per_file.append(len(pdf_temp))
                    except Exception:
                        pages_per_file.append(1)
                total_pages_all = sum(pages_per_file)

            progress_bar = st.progress(0.0)
            
            # Cột 1: File (ĐẦU) | Cột 2: Tiến độ (GIỮA) | Cột 3: Thời gian (CUỐI)
            col_file, col_progress, col_time = st.columns(3)
            file_placeholder = col_file.empty()
            progress_placeholder = col_progress.empty()
            time_placeholder = col_time.empty()
            
            start_time = time.time()
            file_results = []
            tracker = {'pages_done': 0}
            total_extracted_all = 0
            
            for f_idx, up_file in enumerate(uploaded_files):
                file_name = up_file.name
                file_bytes = up_file.read()
                f_pages = pages_per_file[f_idx]
                
                def on_page_done(page_num, total_p):
                    tracker['pages_done'] += 1
                    pages_done = tracker['pages_done']
                    
                    elapsed = time.time() - start_time
                    avg_time = elapsed / pages_done
                    remaining_pages = max(total_pages_all - pages_done, 0)
                    remaining_sec = max(int(remaining_pages * avg_time), 1) if pages_done < total_pages_all else 0
                    
                    pct = min(pages_done / total_pages_all, 1.0)
                    progress_bar.progress(pct)
                    
                    # 1. CỘT ĐẦU TIÊN: File 1/1 và Tên file
                    file_placeholder.markdown(f"""
                    <div class="glow-card-file">
                        <div class="glow-val" style="color: #FBBF24;">File {f_idx + 1}/{total_files}</div>
                        <div class="glow-label">{file_name[:16]}</div>
                    </div>
                    """, unsafe_allow_html=True)
                    
                    # 2. CỘT GIỮA: Tiến độ % và số trang
                    progress_placeholder.markdown(f"""
                    <div class="glow-card-prog">
                        <div class="glow-val" style="color: #34D399;">{int(pct * 100)}%</div>
                        <div class="glow-label">Trang {pages_done}/{total_pages_all}</div>
                    </div>
                    """, unsafe_allow_html=True)
                    
                    # 3. CỘT CUỐI: Thời gian đếm ngược (đủ cả PHÚT và GIÂY)
                    time_placeholder.markdown(f"""
                    <div class="glow-card-time">
                        <div class="glow-val" style="color: #38BDF8;">⏳ {format_time_str(remaining_sec)}</div>
                        <div class="glow-label">Thời gian còn lại</div>
                    </div>
                    """, unsafe_allow_html=True)
                    
                records = extract_from_single_pdf(file_bytes, page_callback=on_page_done)
                file_results.append((file_name, records))
                total_extracted_all += len(records)
                
            total_time_spent = int(time.time() - start_time)
            progress_bar.progress(1.0)
            
            # Tạo file Excel - TÊN FILE ĐỒNG NHẤT: PDF-TO-EXCEL.xlsx
            excel_bytes = create_excel_multi_sheet(file_results)
            out_name = "PDF-TO-EXCEL.xlsx"
            
            opened_path = auto_open_local_file(excel_bytes, out_name)
            
            st.session_state['excel_data'] = excel_bytes
            st.session_state['excel_name'] = out_name
            st.session_state['total_files'] = total_files
            st.session_state['total_extracted'] = total_extracted_all
            st.session_state['total_time_str'] = format_time_str(total_time_spent)
            st.session_state['opened_path'] = opened_path
            st.session_state['completed'] = True
            st.rerun()

# ==================== KHI XỬ LÝ XONG ====================
if st.session_state.get('completed', False):
    out_name = st.session_state['excel_name']
    excel_bytes = st.session_state['excel_data']
    opened_path = st.session_state.get('opened_path', None)
    
    render_auto_downloader(excel_bytes, out_name)
    open_msg = "Đã tự động mở file trên màn hình của bạn!" if opened_path else "Đang tự động tải về thư mục Downloads..."
    
    st.markdown(f"""
    <div class="cyber-success">
        <div style="display: flex; align-items: center; margin-bottom: 6px;">
            <span style="font-size: 24px; margin-right: 10px;">✨</span>
            <h3 style="margin: 0; color: #34D399; font-size: 18px; font-weight: 800;">
                XỬ LÝ THÀNH CÔNG TRONG {st.session_state['total_time_str']}
            </h3>
        </div>
        <div style="font-size: 13px; line-height: 1.6; color: #E2E8F0;">
            • Quét hoàn tất: <b>{st.session_state['total_files']}</b> file PDF &nbsp;|&nbsp; Tìm thấy: <b style="color: #38BDF8;">{st.session_state['total_extracted']}</b> phiếu SM hợp lệ.<br>
            • 📥 <b>Tên file: <code>{out_name}</code></b> ({open_msg})
        </div>
    </div>
    """, unsafe_allow_html=True)
    
    # Nút XỬ LÝ FILE MỚI (Refresh trang về ban đầu)
    if st.button("🔄 XỬ LÝ FILE MỚI (LÀM MỚI TRANG)", type="primary", use_container_width=True):
        for k in ['excel_data', 'excel_name', 'total_files', 'total_extracted', 'total_time_str', 'completed', 'opened_path']:
            if k in st.session_state:
                del st.session_state[k]
        st.rerun()

    st.markdown("""
    <div style='text-align: center; margin-top: 10px; font-size: 12px; color: #94A3B8;'>
        💡 <b>Mẹo tự động mở file trên trình duyệt:</b> Nhấp chuột phải vào file <code>PDF-TO-EXCEL.xlsx</code> vừa tải ở góc trình duyệt ➔ Chọn <i>"Luôn mở các tệp loại này"</i> để máy tự bật Excel mỗi khi tải xong!
    </div>
    """, unsafe_allow_html=True)
    
    st.markdown("<div style='text-align: center; margin-top: 4px;'>", unsafe_allow_html=True)
    st.download_button(
        label="📥 Bấm vào đây nếu muốn tải lại file Excel",
        data=excel_bytes,
        file_name=out_name,
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    st.markdown("</div>", unsafe_allow_html=True)
